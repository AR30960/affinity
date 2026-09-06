import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_document():
    doc = docx.Document()
    
    # Page setup - Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Styles base
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("AFFINITY")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(28)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Dark Blue
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Spécifications et Documentation Fonctionnelle du Moteur d'Affinités Électives\n(Basé sur le projet de M. ROGER André)")
    run_sub.font.size = Pt(14)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
    
    doc.add_paragraph() # Spacer
    
    # Encadré d'introduction
    table_intro = doc.add_table(rows=1, cols=1)
    table_intro.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_intro = table_intro.cell(0, 0)
    set_cell_background(cell_intro, "F0F4F8")
    p_intro = cell_intro.paragraphs[0]
    p_intro.paragraph_format.space_before = Pt(8)
    p_intro.paragraph_format.space_after = Pt(8)
    r_intro = p_intro.add_run("Note importante : Ce document sert de réceptacle partagé entre vous et l'assistant Antigravity. Il décrit le fonctionnement actuel de l'application. Vous pouvez ajouter, modifier ou préciser des exigences directement dans ce fichier. Lorsque vous direz « prends en compte », l'assistant analysera les nouveautés et les appliquera dans la base de données et l'application Web.")
    r_intro.font.size = Pt(10.5)
    r_intro.font.italic = True
    r_intro.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

    doc.add_paragraph() # Spacer

    def add_heading_1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(18)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        return h

    def add_heading_2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
        return h

    # Section 1
    add_heading_1("1. VUE D'ENSEMBLE DE L'APPLICATION AFFINITY")
    p = doc.add_paragraph()
    p.add_run("L'application Affinity est un moteur de calcul d'affinités électives sophistiqué. Elle permet de mesurer le degré de compatibilité (en pourcentage) entre deux individus à partir de leurs profils, cartes d'identité détaillées et leurs réponses à un questionnaire structuré par catégories et thématiques.")

    add_heading_2("Objectifs principaux :")
    doc.add_paragraph("• Gérer des profils utilisateurs avec 3 niveaux de rôles (Administrateur, Abonné, Invité).", style='List Bullet')
    doc.add_paragraph("• Proposer des fiches d'identité riches (caractéristiques physiques, morphologiques, origines, style de vie, ville/géolocalisation).", style='List Bullet')
    doc.add_paragraph("• Organiser les questions par Packs (domaines), par Cible (Homme, Femme, Tous) et par Classe de confidentialité/sensibilité.", style='List Bullet')
    doc.add_paragraph("• Permettre des évaluations par critères de Goût (G) et par critères MULTI (Voulus, Acceptés, Démotivants, Prohibés).", style='List Bullet')
    doc.add_paragraph("• Calculer un score d'affinité bilatéral précis en tenant compte de la distance kilométrique et des orientations sexuelles.", style='List Bullet')
    doc.add_paragraph("• Offrir une interface Web moderne (Dashboard, Questionnaire, Matrice d'affinité, Panneau d'administration).", style='List Bullet')

    # Section 2
    add_heading_1("2. STRUCTURE DES DONNÉES & CLASSES DE QUESTIONS")
    p = doc.add_paragraph("La base de données gère la classification rigoureuse des informations et des droits d'accès :")

    add_heading_2("Classes de sensibilité des questions :")
    t_class = doc.add_table(rows=7, cols=3)
    t_class.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Niveau", "Classe", "Description / Confidentialité"]
    for i, h in enumerate(headers):
        cell = t_class.cell(0, i)
        set_cell_background(cell, "1E3A8A")
        p_h = cell.paragraphs[0]
        r_h = p_h.add_run(h)
        r_h.font.bold = True
        r_h.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    data_classes = [
        ("1", "Standards", "Questions générales sur les centres d'intérêt et habitudes quotidiennes."),
        ("2", "Personnelles", "Questions sur les convictions, le mode de vie et la personnalité."),
        ("3", "Intimes", "Questions abordant la vie privée et les attentes relationnelles approfondies."),
        ("4", "Privées", "Questions confidentielles réservées aux abonnés ayant validé un échange."),
        ("5", "À caractère sexuel", "Questions relatives à la sexualité et aux préférences intimes."),
        ("9", "Interdits / Fantasmes", "Questions très spécifiques sujettes à autorisation explicite.")
    ]
    for row_idx, data in enumerate(data_classes, start=1):
        bg = "F9FAFB" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = t_class.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            cell.paragraphs[0].add_run(text)

    doc.add_paragraph() # Spacer

    add_heading_2("Types de questions :")
    doc.add_paragraph("• Type G (Goûts uniques) : Évaluation d'une préférence simple sur une échelle de 1 à 10.", style='List Bullet')
    doc.add_paragraph("• Type MULTI (V, A, D, P) : Évaluation multidimensionnelle basée sur 4 dimensions :", style='List Bullet')
    doc.add_paragraph("    - V (Voulu) : Ce que la personne recherche prioritairement.", style='List Bullet')
    doc.add_paragraph("    - A (Accepté) : Ce que la personne tolère ou accepte volontiers.", style='List Bullet')
    doc.add_paragraph("    - D (Démotivant) : Ce qui freine ou rebute la personne.", style='List Bullet')
    doc.add_paragraph("    - P (Prohibé) : Ce qui constitue un rédhibitoire absolu.", style='List Bullet')

    # Section 3
    add_heading_1("3. ALGORITHME DE CALCUL D'AFFINITÉ")
    p = doc.add_paragraph("Le score d'affinité globale (entre 0% et 100%) entre un Profil A et un Profil B repose sur :")
    doc.add_paragraph("1. La compatibilité des cartes d'identité (filtre de genre/cible, orientation sexuelle).", style='List Bullet')
    doc.add_paragraph("2. La distance géographique (calcul Haversine basé sur les coordonnées des villes en km).", style='List Bullet')
    doc.add_paragraph("3. La correspondance des réponses au questionnaire pour les questions communes autorisées :", style='List Bullet')
    doc.add_paragraph("   - Pour les questions G : Écart de note ramené à une valeur de proximité.", style='List Bullet')
    doc.add_paragraph("   - Pour les questions MULTI : Analyse du croisement entre ce que A veut/accepte et ce que B propose/est.", style='List Bullet')

    # Section 4
    add_heading_1("4. FONCTIONNALITÉS WEB IMPLEMENTÉES")
    doc.add_paragraph("• Dashboard dynamique avec statistiques en temps réel et filtres avancés.", style='List Bullet')
    doc.add_paragraph("• Gestionnaire de profil (Édition de pseudo, rôle, carte d'identité complète).", style='List Bullet')
    doc.add_paragraph("• Module de questionnaire adaptatif avec filtrage selon les droits et la cible.", style='List Bullet')
    doc.add_paragraph("• Matrice d'affinité bilatérale avec comparateur détaillé question par question.", style='List Bullet')
    doc.add_paragraph("• Panneau d'administration pour valider les demandes d'accès aux classes restreintes.", style='List Bullet')

    # Section 5
    add_heading_1("5. GUIDE DE COLLABORATION (COMMENT PROPOSER DES ÉVOLUTIONS)")
    p_collab = doc.add_paragraph()
    p_collab.add_run("Pour faire évoluer l'application via ce document :\n\n")
    p_collab.add_run("1. Ajoutez ou modifiez le texte directement dans les sections ci-dessus (par exemple : nouvelles règles de calcul, nouvelles classes de questions, nouveaux champs sur la carte d'identité, nouvelles contraintes d'ergonomie).\n")
    p_collab.add_run("2. Vous pouvez surligner, mettre en couleur ou ajouter une section « NOUVELLES DEMANDES ».\n")
    p_collab.add_run("3. Dans la discussion avec l'assistant, dites simplement : « Prends en compte les modifications d'Affinity.docx ».\n")
    p_collab.add_run("4. L'assistant lira le fichier .docx, identifiera automatiquement vos ajouts/modifications et mettra à jour le code source et la base de données de l'application.")

    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Affinity.docx")
    doc.save(output_path)
    print(f"Document créé avec succès : {output_path}")

if __name__ == "__main__":
    create_document()
