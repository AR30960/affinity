import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Helpers de style docx
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_document_language(doc, lang_code='fr-FR'):
    styles_elm = doc.styles.element
    doc_defaults = styles_elm.xpath('w:docDefaults')
    if doc_defaults:
        rPrDefault = doc_defaults[0].find(qn('w:rPrDefault'))
        if rPrDefault is None:
            rPrDefault = OxmlElement('w:rPrDefault')
            doc_defaults[0].append(rPrDefault)
        rPr = rPrDefault.find(qn('w:rPr'))
        if rPr is None:
            rPr = OxmlElement('w:rPr')
            rPrDefault.append(rPr)
        for l in rPr.findall(qn('w:lang')):
            rPr.remove(l)
        lang = OxmlElement('w:lang')
        lang.set(qn('w:val'), lang_code)
        lang.set(qn('w:eastAsia'), lang_code)
        lang.set(qn('w:bidi'), 'ar-SA')
        rPr.append(lang)

def add_header(doc, title, subtitle, meta):
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(title)
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(28)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Bleu nuit

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run(subtitle)
    r_sub.font.size = Pt(13)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = p_meta.add_run(meta)
    r_meta.font.size = Pt(10)
    r_meta.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    doc.add_paragraph()

def add_callout(doc, text, fill_hex="F0F4F8", border_color_rgb=(0x1E, 0x40, 0xAF)):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    set_cell_background(cell, fill_hex)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.right_indent = Inches(0.15)
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    r.font.italic = True
    r.font.color.rgb = RGBColor(*border_color_rgb)
    doc.add_paragraph()

def add_h1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(18)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    return h

def add_h2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x02, 0x84, 0xC7) # Cyan / Bleu
    return h

def add_p(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(11)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    return p

# ==============================================================================
# 1. GÉNÉRATION DU MANUEL UTILISATEUR
# ==============================================================================
def generate_user_manual():
    doc = docx.Document()
    set_document_language(doc, 'fr-FR')

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Titre et Entête
    add_header(doc, 
               "AFFINITY",
               "Manuel d'Utilisation Officiel — Moteur d'Affinités Électives\nGuide Pratique, Éthique Relationnelle & Parcours Membre",
               "Auteur : M. André ROGER  |  Version 2.0 (Mise à jour Septembre 2026)  |  Application Web & Mobile")

    add_callout(doc, 
                "« Affinity n'est pas un catalogue de visages ni une roulette de rencontres éphémères. C'est un instrument d'élection réciproque et d'accord profond entre deux êtres humains qui choisissent de se découvrir dans la vérité de ce qu'ils ont vécu, de ce qu'ils vivent aujourd'hui, et de ce qu'ils désirent bâtir ensemble. »\n— Conception & Vision par M. André ROGER",
                "F8FAFC", (0x0F, 0x17, 0x2A))

    # Sommaire
    add_h2(doc, "TABLE DES MATIÈRES")
    toc_items = [
        "1. Présentation d'Affinity : « À la recherche de » — Philosophie & Démarche Élective",
        "2. Mon Profil & Fiche d'Identité : Renseignement et Complétude Obligatoire",
        "3. Les Modules d'Identité : « + sur moi » et « + sur l'autre » (Classe 8)",
        "4. Les Questionnaires Thématiques et la Matrice Canonique Type M (V-A-D-P)",
        "5. Les Jeux de Questions : Packs d'Appartenance (Jeu 1, Jeu 2, Jeu 3)",
        "6. Le Centre de Matchs : Demande Réciproque, Négociation du Périmètre & Restitution",
        "7. L'Historique Bilatéral des Matchs & Consultation des Rapports d'Affinité",
        "8. Gestion des Droits d'Accès aux Classes & Passage au Statut Abonné",
        "9. Gestion du Compte, Sécurité du Mot de Passe & Foire Aux Questions (FAQ)"
    ]
    for item in toc_items:
        add_bullet(doc, item)
    doc.add_paragraph()

    # SECTION 1 : À LA RECHERCHE DE
    add_h1(doc, "1. PRÉSENTATION D'AFFINITY : « À LA RECHERCHE DE »")
    add_p(doc, "Bienvenue sur Affinity. Si vous utilisez cette application, c'est que vous refusez la superficialité des sites de rencontre traditionnels. Sur la plupart des plateformes, la mise en relation repose sur une image trompeuse, un geste réflexe (« swiper » à droite ou à gauche) et un échange superficiel qui débouche trop souvent sur la désillusion. Affinity prend le contre-pied absolu de cette approche en fondant toute sa démarche sur le principe des Affinités Électives.")
    
    add_h2(doc, "1.1 Le Cœur de la Démarche : « À la recherche de »")
    add_p(doc, "Dans votre profil Affinity, la rubrique « À la recherche de » n'est pas une simple case administrative : c'est le cap fondateur de votre présence parmi nous. Elle s'articule autour de quatre quêtes fondamentales :")

    add_bullet(doc, "Avant de pouvoir rencontrer l'autre de manière constructive, il est nécessaire de savoir avec lucidité qui l'on est. Grâce aux questionnaires thématiques et au module d'identité « + sur moi », Affinity vous invite à poser un regard honnête sur votre histoire : vos passions réelles, votre rythme quotidien, vos goûts affirmés, vos valeurs morales et votre style de vie.", "1. À la recherche de soi-même : ")

    add_bullet(doc, "Combien de relations échouent parce que les attentes mutuelles étaient tues ou idéalisées ? Dans le module miroir « + sur l'autre », vous explicitez ce que vous désirez partager, ce qui vous émerveille chez un partenaire, mais aussi ce qui vous est inconfortable ou inacceptable (cadre de vie, habitudes tabac/alcool, pratique sportive, vision de l'engagement). Vous pouvez également déclarer une « Indifférence bienveillante » sur les critères qui ne constituent aucun obstacle pour vous.", "2. À la recherche de l'autre dans sa singularité : ")

    add_bullet(doc, "L'amour et l'amitié véritable ne naissent pas de la ressemblance parfaite, mais de la résonance. Affinity ne recherche pas votre double identique, mais votre harmonie dynamique. L'algorithme analyse comment votre désir de découverte (votre appétence pour de nouvelles expériences) rencontre la capacité et l'envie de partage de l'autre, révélant ainsi vos précieux « Points de Fusion » et signalant avec bienveillance vos « Zones de Vigilance ».", "3. À la recherche d'une synergie dynamique : ")

    add_bullet(doc, "Sur Affinity, rien n'est imposé, tout est consenti. Aucune information intime ou sensible n'est livrée au hasard. Chaque mise en relation (Match) fait l'objet d'un accord préalable bilatéral où chaque partenaire choisit son niveau de confort : partager uniquement des pourcentages globaux par thématique, ou accepter d'ouvrir le détail question par question. Vous avancez au rythme de votre confiance mutuelle.", "4. À la recherche d'une relation respectueuse et consentie : ")

    # SECTION 2 : MON PROFIL
    add_h1(doc, "2. MON PROFIL & FICHE D'IDENTITÉ OBLIGATOIRE")
    add_p(doc, "L'onglet « Mon Profil » constitue votre cockpit personnel. Conformément aux règles fondamentales conçues par M. André ROGER, chaque profil doit obligatoirement avoir complété l'ensemble de ses informations essentielles pour prétendre au calcul d'affinité avec les autres membres.")
    add_bullet(doc, "Votre identifiant public visible par la communauté.", "• Pseudo public : ")
    add_bullet(doc, "Préservés avec une stricte confidentialité par le système.", "• Nom et Prénom d'état civil : ")
    add_bullet(doc, "Homme (1) ou Femme (2). Conditionne les questions ciblées et les recherches réciproques.", "• Sexe biologique : ")
    add_bullet(doc, "Permet le calcul automatique de votre âge et de votre tranche de maturité.", "• Date de naissance : ")
    add_bullet(doc, "Pays, département français (de 01 à 976) et commune. Détermine le calcul de distance kilométrique selon la formule de Haversine.", "• Localisation géographique (« J'habite ici ») : ")
    add_bullet(doc, "Échanges et Amitié, Recherche de l'Âme Sœur, Relation Sérieuse & Durable, Partage de Passions, etc.", "• À la recherche de : ")
    add_bullet(doc, "Un texte sincère décrivant vos aspirations, votre univers et votre vision du bonheur partagé.", "• Présentation générale : ")
    add_callout(doc, "Règle d'or de la complétude : La jauge de complétude de votre profil doit atteindre 100%. Tant que votre profil ou vos modules d'identité ne sont pas entièrement renseignés, vous n'apparaîtrez pas dans la liste des profils complétés disponibles et ne pourrez pas initier de demande de match.", "FEF3C7", (0x92, 0x40, 0x0E))

    # SECTION 3 : CLASSE 8 IDENTITÉ
    add_h1(doc, "3. LES MODULES D'IDENTITÉ : « + SUR MOI » ET « + SUR L'AUTRE »")
    add_p(doc, "En haut à droite de votre profil se trouvent deux boutons d'accès majeurs consacrés à la Classe 8 (Identité) :")
    add_h2(doc, "3.1 Module « + sur moi » (Type P - Caractéristiques personnelles)")
    add_p(doc, "Vous y renseignez vos réalités morphologiques, esthétiques et d'hygiène de vie : taille (stature en cm), poids (kg), silhouette, pointure, couleur des yeux et des cheveux, pilosité faciale (pour les hommes), bonnet et tour de poitrine (pour les femmes), style vestimentaire, origines culturelles, rythme de vie et cadre d'épanouissement. Récemment enrichi, ce module intègre également vos habitudes vis-à-vis du tabac/vape, de l'alcool, votre régime alimentaire, votre pratique sportive, votre cohabitation avec des animaux de compagnie et le port de tatouages/piercings.")
    
    add_h2(doc, "3.2 Module « + sur l'autre » (Type T - Critères de tolérance et préférences)")
    add_p(doc, "Dans ce module miroir, vous délimitez vos préférences chez le partenaire recherché : plages numériques acceptées (ex : taille entre 165 cm et 185 cm) ou choix multiples de silhouettes, couleurs, styles et modes de vie. Si un aspect vous importe peu, cochez simplement la case « Indifférent » : elle accordera automatiquement une compatibilité maximale sur ce critère sans exiger de contrainte inutile.")

    # SECTION 4 : QUESTIONNAIRES ET TYPE M
    add_h1(doc, "4. LES QUESTIONNAIRES THÉMATIQUES ET LE TYPE M")
    add_p(doc, "L'onglet « Questionnaires » regroupe l'ensemble des questions réparties en classes de sensibilité croissante : Classe 1 (Standards), Classe 2 (Personnelles), Classe 3 (Intimes), Classe 4 (Privées), Classe 5 (Hors normes / À caractère sexuel) et Classe 9 (Amorales / Interdits).")
    add_h2(doc, "4.1 Le Type G (Goûts simples)")
    add_p(doc, "Évaluation directe de votre appréciation sur une échelle de 1 (Faible) à 5 (Très fort), ou 9 (Pas du tout / Jamais).")
    add_h2(doc, "4.2 Le Type M (La matrice canonique à 4 axes)")
    add_p(doc, "Véritable signature d'Affinity, le Type M décompose chaque sujet selon quatre dimensions complémentaires :")
    add_bullet(doc, "Votre expérience passée. Avez-vous déjà pratiqué ou expérimenté ce sujet ? (Note 1 à 5, ou 9 si jamais).", "• Axe V (Le Vécu) : ")
    add_bullet(doc, "Votre pratique présente. Ce sujet fait-il partie de vos habitudes actuelles ?", "• Axe A (Actuel / Présent) : ")
    add_bullet(doc, "Votre désir futur. Souhaitez-vous découvrir ou poursuivre cette pratique à l'avenir ?", "• Axe D (Découverte / Futur) : ")
    add_bullet(doc, "Votre tolérance au partage. Acceptez-vous que votre partenaire s'y adonne ou le pratique avec vous ?", "• Axe P (Partage) : ")

    # SECTION 5 : LES PACKS / JEUX
    add_h1(doc, "5. LES JEUX D'APPARTENANCE (JEU 1, JEU 2, JEU 3)")
    add_p(doc, "Dans toutes les vues de l'application (Questionnaire membre, Banque de questions, Modération et Restitution détaillée), chaque question arbore désormais un badge distinctif indiquant son jeu d'appartenance :")
    add_bullet(doc, "Le tronc commun fondamental des questions mères standards et d'identité, accessible à l'ensemble des utilisateurs.", "• 🎮 Jeu 1 : ")
    add_bullet(doc, "Les questions de précision attachées aux questions d'origine, permettant d'approfondir un domaine spécifique.", "• 🎮 Jeu 2 : ")
    add_bullet(doc, "Les thématiques avancées, intimes et sensuelles réservées aux abonnés confirmés.", "• 🎮 Jeu 3 : ")

    # SECTION 6 : LE CENTRE DE MATCHS
    add_h1(doc, "6. LE CENTRE DE MATCHS & NÉGOCIATION DU PÉRIMÈTRE")
    add_p(doc, "L'accès au calcul de compatibilité s'effectue depuis l'onglet « Centre de Matchs » (ou « Diagnostic & Matchs »). L'ancien bouton instantané du bandeau supérieur a été supprimé afin de sanctuariser un processus rigoureux et bilatéral :")
    add_h2(doc, "6.1 Sélection d'un profil disponible complété")
    add_p(doc, "La liste affiche uniquement les profils du sexe opposé (ou correspondant à votre recherche) ayant atteint 100% de complétude sur leur profil et leurs modules d'identité. Vous y visualisez leur pseudo, ville, âge, distance et profession.")
    add_h2(doc, "6.2 Personnalisation du niveau de restitution par classe")
    add_p(doc, "Avant d'envoyer votre demande de match, vous déterminez le périmètre de classes souhaité et positionnez individuellement pour chaque classe le niveau de transparence désiré :")
    add_bullet(doc, "Restitution synthétique sous forme de pourcentages d'affinité par thématique et par sujet. Idéal pour préserver sa pudeur tout en évaluant la concordance globale.", "• 📊 Niveau Pourcentage (%) : ")
    add_bullet(doc, "Confrontation transparente question par question des réponses respectives (avec analyse croisée des axes). Réservé aux thématiques où les deux membres souhaitent une clarté totale.", "• 🔍 Niveau Détail : ")
    add_h2(doc, "6.3 Réception et arbitrage réciproque")
    add_p(doc, "Le destinataire reçoit la demande dans sa boîte de réception des matchs. Il peut accepter les conditions proposées ou ajuster le niveau de restitution (par exemple basculer une classe de « Détail » vers « % »). Le calcul final du moteur n'autorisera le détail question par question que sur les classes ayant reçu un accord bilatéral explicite.")

    # SECTION 7 : HISTORIQUE BILATÉRAL
    add_h1(doc, "7. L'HISTORIQUE BILATÉRAL & FENÊTRE DE RÉSULTATS DÉDIÉE")
    add_p(doc, "Le résultat d'un match ne s'affiche plus de manière fugitive sur la page courante : il est archivé de manière permanente dans votre Historique des Matchs.")
    add_p(doc, "En cliquant sur n'importe quelle ligne de match réalisé ou sur le bouton « 🔍 Détail », une fenêtre modale plein écran dédiée s'ouvre, comprenant :")
    add_bullet(doc, "Une jauge circulaire lumineuse calculant l'affinité mathématique globale pondérée selon la sensibilité des classes.", "• Score Global d'Affinité : ")
    add_bullet(doc, "Indicateurs visuels de concordance sur le Vécu (V), le Présent (A), les Goûts (G) et la Synergie Croisée (D ⇄ P).", "• Barres d'Équilibre par Axe : ")
    add_bullet(doc, "Vérification bilatérale des critères physiques et de mode de vie (+ sur moi vs + sur l'autre), avec verdict immédiat (Compatible, Partiel, Non compatible).", "• Diagnostic Préalable d'Identité : ")
    add_bullet(doc, "Mise en lumière des thématiques où votre accord dépasse 85%, et alertes prévenantes lorsque l'accord descend sous 35%.", "• Points de Fusion & Zones de Vigilance : ")
    add_bullet(doc, "Pour les classes convenues en mode détail, confrontation comparative question par question avec badge de jeu et score unitaire.", "• Confrontation Détaillée Question par Question : ")

    # SECTION 8 : ACCÈS ET ABONNEMENT
    add_h1(doc, "8. GESTION DES DROITS D'ACCÈS & STATUT ABONNÉ")
    add_p(doc, "Affinity propose deux niveaux de privilèges :")
    add_bullet(doc, "Accès gratuit au Jeu 1 et à la Classe 1. Permet de compléter son profil, découvrir les questionnaires et tester les mécanismes de base.", "• Statut Invité : ")
    add_bullet(doc, "Accès étendu aux classes supérieures (Classes 2, 3, 4, 5 et 9) et aux Jeux avancés. Possibilité de formuler des demandes d'accès d'un clic depuis votre cockpit.", "• Statut Abonné : ")
    add_p(doc, "L'administrateur valide désormais les demandes d'accès aux périmètres en un temps record grâce à un tableau de modération unifié et un bouton « Tout accorder » par lot.")

    # SECTION 9 : FAQ
    add_h1(doc, "9. GESTION DU COMPTE, SÉCURITÉ & FAQ")
    add_bullet(doc, "Accessible directement depuis l'onglet Mon Profil via le formulaire sécurisé avec salage cryptographique.", "• Modifier son mot de passe : ")
    add_bullet(doc, "Vos réponses et fiches d'identité sont stockées localement et protégées. Aucune donnée n'est vendue ni cédée à des tiers publicitaires.", "• Confidentialité absolue : ")
    add_bullet(doc, "Vérifiez que vous avez atteint 100% sur Mon Profil, « + sur moi » et « + sur l'autre ». L'autre personne doit également avoir atteint 100%.", "• Pourquoi un membre n'apparaît pas dans la liste des matchs ? ")

    output_path = "Manuel_Utilisateur_Affinity.docx"
    doc.save(output_path)
    print(f"Manuel Utilisateur généré avec succès : {output_path}")

# ==============================================================================
# 2. GÉNÉRATION DU MANUEL TECHNIQUE
# ==============================================================================
def generate_technical_manual():
    doc = docx.Document()
    set_document_language(doc, 'fr-FR')

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Titre et Entête
    add_header(doc, 
               "AFFINITY — Manuel Technique",
               "Dossier d'Architecture Logicielle, Moteur d'Affinité Élective & Guide d'Exploitation",
               "Référence Projet : ARP001  |  Auteur : M. André ROGER  |  Version 2.0 (Septembre 2026)")

    add_callout(doc, 
                "Document technique de référence à destination des développeurs, administrateurs système et exploitants de la plateforme Affinity. Ce document détaille l'architecture logicielle, les mécanismes de concurrence SQLite WAL, la formalisation mathématique des algorithmes V-A-D-P et les spécifications de l'API REST.",
                "F1F5F9", (0x1E, 0x29, 0x3B))

    # Sommaire
    add_h2(doc, "TABLE DES MATIÈRES TECHNIQUE")
    toc_items = [
        "1. Vue d'Ensemble de l'Architecture & Stack Technologique",
        "2. Haute Concurrence SQLite & Mode Journal WAL (Windows / Multi-Threads)",
        "3. Modèle de Données & Schéma Relationnel Complet (affinity.db)",
        "4. Moteur d'Identité & Évaluation Croisée Classe 8 (evaluate_identity_compatibility)",
        "5. Algorithme du Moteur d'Affinité Canonique (calculate_affinity & V-A-D-P)",
        "6. Négociation Réciproque du Niveau de Restitution par Classe (% vs Détail)",
        "7. Supervision Administrateur & Matchs Discrets Confidentiels",
        "8. Spécification Complète des API REST (Endpoints HTTP GET, POST, PUT, DELETE)",
        "9. Sauvegardes Locales, Sécurité & Déploiement Continu (Render / Windows 10 & 11)"
    ]
    for item in toc_items:
        add_bullet(doc, item)
    doc.add_paragraph()

    # SECTION 1 : ARCHITECTURE
    add_h1(doc, "1. VUE D'ENSEMBLE DE L'ARCHITECTURE & CHOIX TECHNIQUES")
    add_p(doc, "Affinity (ARP001) est une application hybride Web & Locale ultra-légère conçue sans frameworks lourds pour maximiser les performances, l'autonomie et la pérennité :")
    add_bullet(doc, "HTML5 sémantique, CSS3 Vanilla moderne (variables, glassmorphism, flexbox/grid responsive) et JavaScript ES6+ pur sans dépendance externe npm.", "• Frontend SPA (Single Page Application) : ")
    add_bullet(doc, "Serveur HTTP multithreadé basé sur le module natif Python socketserver.ThreadingMixIn et http.server, garantissant une exécution instantanée sur Windows 10/11 sans conteneurisation obligatoire.", "• Backend Python 3 Haute Performance : ")
    add_bullet(doc, "Base relationnelle SQLite3 autonome (affinity.db) hautement optimisée, gérée directement par requêtes SQL paramétrées sans surcouche ORM opaque.", "• Moteur de Persistance SQLite3 : ")

    # SECTION 2 : SQLITE WAL
    add_h1(doc, "2. HAUTE CONCURRENCE SQLITE & MODE JOURNAL WAL")
    add_p(doc, "Face aux contraintes multi-threads sous Windows (où les opérations concurrentes d'écriture pouvaient provoquer l'erreur database is locked), l'architecture de données a été renforcée :")
    add_bullet(doc, "Chaque connexion ouverte par get_db() applique systématiquement un timeout étendu de 30.0 secondes (sqlite3.connect(DB_PATH, timeout=30.0)).", "• Timeout de Verrouillage Dédié : ")
    add_bullet(doc, "Exécution de PRAGMA busy_timeout = 30000 ordonnant au moteur SQLite de réitérer automatiquement ses requêtes d'écriture jusqu'à libération du verrou.", "• Pragma Busy Timeout : ")
    add_bullet(doc, "Activation au démarrage dans init_db() du mode Write-Ahead Logging (PRAGMA journal_mode = WAL). Ce mode permet aux lecteurs de lire la base sans bloquer les écrivains, et aux écrivains d'insérer des données sans interrompre les consultations.", "• Mode Journal WAL (Write-Ahead Logging) : ")

    # SECTION 3 : SCHÉMA RELATIONNEL
    add_h1(doc, "3. MODÈLE DE DONNÉES & SCHÉMA RELATIONNEL (affinity.db)")
    add_p(doc, "Le schéma relationnel structure l'ensemble de l'écosystème :")
    add_bullet(doc, "id (PK AUTO), pseudo (UNIQUE), code_profil (UNIQUE), avatar, role ('admin', 'subscriber', 'guest'), password_hash (SHA-256), salt, created_at.", "• profiles : ")
    add_bullet(doc, "profile_id (PK, FK), nom, prenom, sexe (1: Homme, 2: Femme), date_naissance, ville, habite_pays, habite_region_dept, habite_commune, travail_commune, taille, poids, pointure, tour_poitrine, tour_taille, tour_hanches, origines, couleur_cheveux, style, situation_famille, recherche_de, bio.", "• identity_cards : ")
    add_bullet(doc, "id (PK), pack_id (FK), cible (0: Mixte, 1: Homme, 2: Femme), classe (0, 1, 2, 3, 4, 5, 8, 9), thematique, sujet, type ('G', 'M', 'P', 'T'), texte, config_reponses (JSON), status ('validated', 'pending_review'), n_quest_lie (lien vers question mère ou question miroir).", "• questions : ")
    add_bullet(doc, "profile_id, question_id, axis ('G', 'V', 'A', 'D', 'P'), value (INTEGER 1..5 ou 9), updated_at.", "• answers : ")
    add_bullet(doc, "profile_id, question_id, valeur_num (REAL), valeur_text (TEXT), updated_at.", "• identity_answers_self (Module + sur moi) : ")
    add_bullet(doc, "profile_id, question_id, min_val (REAL), max_val (REAL), options_json (TEXT JSON), indifferent (0/1), updated_at.", "• identity_answers_partner (Module + sur l'autre) : ")
    add_bullet(doc, "id (PK AUTO), sender_id (FK), receiver_id (FK), status ('pending', 'accepted', 'rejected'), allowed_classes (JSON), restitution_mode (TEXT ou JSON dictionnaire), is_discreet (0/1), created_at, responded_at.", "• match_requests : ")
    add_bullet(doc, "id (PK AUTO), profile_id (FK), target_type ('classe', 'pack'), target_value, action_type ('grant', 'revoke'), status ('pending', 'approved', 'rejected'), created_at, responded_at.", "• admin_access_requests : ")
    add_bullet(doc, "profile_id (PK, FK), allowed_classes (JSON list d'entiers), allowed_packs (JSON list d'entiers), updated_at.", "• profile_question_access : ")

    # SECTION 4 : MOTEUR IDENTITÉ CLASSE 8
    add_h1(doc, "4. MOTEUR D'IDENTITÉ & ÉVALUATION CROISÉE CLASSE 8")
    add_p(doc, "La fonction evaluate_identity_compatibility(profile1_id, profile2_id, conn) opère une vérification stricte en miroir :")
    add_bullet(doc, "Vérifie si les caractéristiques déclarées par P2 (identity_answers_self de P2) satisfont les exigences définies par P1 (identity_answers_partner de P1).", "• Sens 1 (Attentes de P1 envers P2) : ")
    add_bullet(doc, "Vérifie symétriquement si les caractéristiques de P1 satisfont les attentes de P2.", "• Sens 2 (Attentes de P2 envers P1) : ")
    add_bullet(doc, "Pour les dimensions numériques (mode: numeric : taille, poids, pointure, tour de poitrine), vérifie min_val <= valeur <= max_val.", "• Évaluation Numérique : ")
    add_bullet(doc, "Pour les sélections qualitatives (mode: select : silhouette, yeux, cheveux, barbe, tabac, alcool, alimentation, sport, animaux, tatouages), vérifie la présence de la valeur dans la liste des options acceptées.", "• Évaluation Qualitative : ")
    add_bullet(doc, "Si le critère comporte indifferent = 1, la dimension est immédiatement déclarée COMPATIBLE.", "• Règle de l'Indifférence : ")
    add_bullet(doc, "Retourne COMPATIBLE (100% de concordance), PARTIEL (score >= 60%), ou INCOMPATIBLE.", "• Niveaux de Synthèse : ")

    # SECTION 5 : ALGORITHME V-A-D-P
    add_h1(doc, "5. ALGORITHME DU MOTEUR D'AFFINITÉ (V-A-D-P)")
    add_p(doc, "La fonction calculate_affinity(p1_id, p2_id, allowed_classes, mode_restitution) réalise le calcul mathématique complet :")
    add_bullet(doc, "Calculée sur l'écart absolu |v1 - v2| ramené sur [0, 1]. Le code 9 (Pas du tout / Jamais) est normalisé à 0.0.", "• Similarité Type G : ")
    add_bullet(doc, "Similarité sur l'expérience vécue dans le passé.", "• Axe V (Vécu) : ")
    add_bullet(doc, "Similarité sur les habitudes actuelles partagées.", "• Axe A (Actuel) : ")
    add_bullet(doc, "Croise le désir de découverte de P1 avec l'acceptation de partage de P2 : sim(P1_D, P2_P) et sim(P2_D, P1_P).", "• Synergie Croisée D ⇄ P : ")
    add_bullet(doc, "Les classes intimes et privées disposent de coefficients de pondération supérieurs (C1: 1.0, C2: 1.2, C3: 1.5, C4: 2.0, C5: 2.5, C8: 1.8, C9: 3.0).", "• Pondération par Classe : ")
    add_bullet(doc, "Score >= 85% = Point de Fusion. Score <= 35% = Zone de Vigilance.", "• Détection Heuristique : ")

    # SECTION 6 : NÉGOCIATION PAR CLASSE
    add_h1(doc, "6. NÉGOCIATION PAR CLASSE (% OU DÉTAIL)")
    add_p(doc, "Le champ restitution_mode de la table match_requests stocke un objet JSON sérialisé sous la forme :")
    add_callout(doc, '{\n  "1": "percentage",\n  "2": "detail",\n  "3": "percentage",\n  "4": "percentage",\n  "5": "detail",\n  "9": "percentage"\n}', "F8FAFC", (0x02, 0x84, 0xC7))
    add_p(doc, "Le moteur n'inclut dans le tableau final questions_details que les questions appartenant aux classes pour lesquelles l'accord réciproque est strictement égal à 'detail'. Pour les autres classes, seuls les pourcentages agrégés par thématique et par sujet sont délivrés.")

    # SECTION 7 : MATCHS DISCRETS
    add_h1(doc, "7. SUPERVISION & MATCHS DISCRETS ADMINISTRATEUR")
    add_p(doc, "L'administrateur peut ordonner un match discret (is_discreet = 1) entre deux membres quelconques sans qu'aucune notification ne leur soit transmise. Le calcul est opéré immédiatement avec restitution intégrale (% et détail sur l'ensemble des classes communes) et archivé dans la table match_requests avec le drapeau confidentiel, permettant une inspection immédiate depuis l'écran de supervision administrative.")

    # SECTION 8 : API REST
    add_h1(doc, "8. SPÉCIFICATION COMPLÈTE DES API REST")
    endpoints = [
        ("GET /api/profiles", "Liste des profils complétés avec statut et localisation"),
        ("GET /api/questions", "Liste des questions validées (avec pack_id et pack_nom)"),
        ("GET /api/admin/questions/pending", "Liste des questions en attente de modération"),
        ("POST /api/admin/questions/{id}/validate", "Valide et active une question en attente"),
        ("POST /api/match-requests", "Émet une demande de match avec paramétrage par classe"),
        ("PUT /api/match-requests/{id}/respond", "Accepte ou refuse une demande avec ajustement du niveau"),
        ("GET /api/match-requests/history", "Historique bilatéral des matchs entre deux profils"),
        ("POST /api/match", "Calcule et restitue le rapport d'affinité selon le niveau convenu"),
        ("PUT /api/access-requests/{id}/respond", "Approuve ou refuse unitairement une demande d'accès"),
        ("PUT /api/admin/access-requests/batch-respond", "Approuve ou refuse par lot toutes les demandes d'accès"),
        ("GET /api/admin/backups", "Liste des sauvegardes locales horodatées disponibles"),
        ("POST /api/admin/backups/create", "Déclenche une sauvegarde SQLite immédiate dans backups/")
    ]
    for ep, desc in endpoints:
        add_bullet(doc, desc, f"{ep} : ")

    # SECTION 9 : SAUVEGARDES ET DÉPLOIEMENT
    add_h1(doc, "9. SAUVEGARDES, SÉCURITÉ & DÉPLOIEMENT")
    add_bullet(doc, "Au démarrage de server.py, une copie horodatée affinity_auto_YYYYMMDD_HHMMSS.db est générée dans backups/.", "• Sauvegarde Automatique : ")
    add_bullet(doc, "Double-clic sur Affinity.bat (lance le serveur en arrière-plan et ouvre l'interface sur http://localhost:8765).", "• Exploitation Windows Locale : ")
    add_bullet(doc, "Compatible déploiement conteneurisé Docker et PaaS Cloud (Render) avec détection automatique des disques persistants montés.", "• Déploiement Cloud : ")

    output_path = "Manuel_Technique_Affinity.docx"
    doc.save(output_path)
    print(f"Manuel Technique généré avec succès : {output_path}")

if __name__ == "__main__":
    generate_user_manual()
    generate_technical_manual()
