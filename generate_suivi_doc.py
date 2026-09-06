import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_document_language(doc, lang_code='fr-FR'):
    """Configure la langue du document Word pour le correcteur orthographique (ex: fr-FR)."""
    styles_elm = doc.styles.element
    doc_defaults = styles_elm.xpath('w:docDefaults')
    if doc_defaults:
        rPrDefault = doc_defaults[0].find(qn('w:rPrDefault'))
        if rPrDefault is None:
            rPrDefault = OxmlElement('w:rPrDefault')
            doc_defaults[0].append(rPrDefault)
        rPr = rPrDefault.find(qn('w:rPr'))
        if rPr is None:
            rPr = OxmlElement('w:rPr')
            rPrDefault.append(rPr)
        for l in rPr.findall(qn('w:lang')):
            rPr.remove(l)
        lang = OxmlElement('w:lang')
        lang.set(qn('w:val'), lang_code)
        lang.set(qn('w:eastAsia'), lang_code)
        lang.set(qn('w:bidi'), 'ar-SA')
        rPr.append(lang)

    for s in doc.styles:
        if hasattr(s, 'element') and s.element is not None:
            rPr = s.element.get_or_add_rPr()
            for l in rPr.findall(qn('w:lang')):
                rPr.remove(l)
            lang = OxmlElement('w:lang')
            lang.set(qn('w:val'), lang_code)
            lang.set(qn('w:eastAsia'), lang_code)
            lang.set(qn('w:bidi'), 'ar-SA')
            rPr.append(lang)

    for p in doc.paragraphs:
        for r in p.runs:
            rPr = r._element.get_or_add_rPr()
            for l in rPr.findall(qn('w:lang')):
                rPr.remove(l)
            lang = OxmlElement('w:lang')
            lang.set(qn('w:val'), lang_code)
            lang.set(qn('w:eastAsia'), lang_code)
            lang.set(qn('w:bidi'), 'ar-SA')
            rPr.append(lang)

    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        rPr = r._element.get_or_add_rPr()
                        for l in rPr.findall(qn('w:lang')):
                            rPr.remove(l)
                        lang = OxmlElement('w:lang')
                        lang.set(qn('w:val'), lang_code)
                        lang.set(qn('w:eastAsia'), lang_code)
                        lang.set(qn('w:bidi'), 'ar-SA')
                        rPr.append(lang)

def create_suivi_document():
    doc = docx.Document()
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("AFFINITY - Suivi ARP001")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(28)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Spécifications et Documentation Fonctionnelle du Moteur d'Affinités Électives\n(Document de Suivi Partagé)")
    run_sub.font.size = Pt(14)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
    
    doc.add_paragraph()

    # Intro box
    table_intro = doc.add_table(rows=1, cols=1)
    table_intro.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_intro = table_intro.cell(0, 0)
    set_cell_background(cell_intro, "F0F4F8")
    p_intro = cell_intro.paragraphs[0]
    p_intro.paragraph_format.space_before = Pt(8)
    p_intro.paragraph_format.space_after = Pt(8)
    r_intro = p_intro.add_run("Note importante : Ce document sert de réceptacle partagé entre vous (Anji, ADM-1) et l'assistant Antigravity. Vous pouvez y ajouter ou modifier des demandes. Lorsque vous direz « Prends en compte », l'assistant analysera les nouveautés, les appliquera et marquera les demandes traitées avec « Fait - [Date] » et le texte barré.")
    r_intro.font.size = Pt(10.5)
    r_intro.font.italic = True
    r_intro.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

    doc.add_paragraph()

    def add_heading_1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(18)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        return h

    def add_heading_2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
        return h

    # Section 1
    add_heading_1("1. VUE D'ENSEMBLE DE L'APPLICATION AFFINITY")
    p = doc.add_paragraph()
    p.add_run("L'application Affinity est un moteur de calcul d'affinités électives sophistiqué. Elle mesure le degré de compatibilité (en %) entre deux individus à partir de leurs profils, cartes d'identité détaillées et leurs réponses à un questionnaire structuré.")

    add_heading_2("Objectifs principaux :")
    doc.add_paragraph("• Gérer des profils utilisateurs avec 3 niveaux de rôles (Administrateur, Abonné, Invité) et un code profil unique (ex: ADM-1 pour Anji, AFF-2...).", style='List Bullet')
    doc.add_paragraph("• Proposer des fiches d'identité riches et des questions de niveau 0 / classe Identité rattachées à la personne.", style='List Bullet')
    doc.add_paragraph("• Organiser les questions par Packs, Cibles (Homme, Femme, Mixte/Tous) et Classes (0-Non définies, 1-Standards, 2-Personnelles, 3-Intimes, 4-Privées, 5-A caractère sexuel, 8-Identité, 9-Interdits/Fantasmes).", style='List Bullet')
    doc.add_paragraph("• Permettre des évaluations par type Goût (G), par type M (Vécu, Actuel, Découverte/Poursuite, Partage), Type P (+ sur vous) et Type T (+ sur l'autre).", style='List Bullet')
    doc.add_paragraph("• Calculer un score d'affinité bilatéral (la distance kilométrique restant purement informative).", style='List Bullet')

    # Section 2
    add_heading_1("2. STRUCTURE DES DONNÉES & CLASSES DE QUESTIONS")
    
    add_heading_2("Classes de sensibilité des questions :")
    t_class = doc.add_table(rows=9, cols=3)
    t_class.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Niveau", "Classe", "Description / Confidentialité"]
    for i, h in enumerate(headers):
        cell = t_class.cell(0, i)
        set_cell_background(cell, "1E3A8A")
        p_h = cell.paragraphs[0]
        r_h = p_h.add_run(h)
        r_h.font.bold = True
        r_h.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    data_classes = [
        ("0", "Non définies", "Questions non définies d'origine dans la base de données."),
        ("1", "Standards", "Questions générales (seule classe accessible par défaut aux Invités)."),
        ("2", "Personnelles", "Questions sur les convictions, le mode de vie et la personnalité."),
        ("3", "Intimes", "Questions abordant la vie privée et les attentes relationnelles."),
        ("4", "Privées", "Questions confidentielles réservées aux abonnés ayant validé un échange."),
        ("5", "A caractère sexuel", "Questions relatives à la sexualité et aux préférences intimes."),
        ("8", "Identité", "Caractéristiques personnelles précises (+ sur vous) et critères de tolérance pour l'autre (+ sur l'autre)."),
        ("9", "Interdits / Fantasmes", "Questions très spécifiques sujettes à autorisation explicite.")
    ]
    for row_idx, data in enumerate(data_classes, start=1):
        bg = "F9FAFB" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = t_class.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            cell.paragraphs[0].add_run(text)

    doc.add_paragraph()

    add_heading_2("Types de questions :")
    doc.add_paragraph("• Type G (Goûts uniques) : Évaluation d'une préférence simple sur une échelle de 1 à 10.", style='List Bullet')
    doc.add_paragraph("• Type M (Multi-Axes) : Évaluation multidimensionnelle basée sur 4 dimensions :", style='List Bullet')
    doc.add_paragraph("    - (V) Le Vécu (Passé)", style='List Bullet')
    doc.add_paragraph("    - (A) Actuel (Présent)", style='List Bullet')
    doc.add_paragraph("    - (D) Découverte ou Poursuite (Futur)", style='List Bullet')
    doc.add_paragraph("    - (P) Partage (Chez la personne qui partage votre quotidien ou chez les autres)", style='List Bullet')
    doc.add_paragraph("• Type P (Précis) : Caractéristiques personnelles précises de classe 8 (Identité) (renseignées dans le module « + sur moi »).", style='List Bullet')
    doc.add_paragraph("• Type T (Tolérance) : Critères de tolérance, limites numériques (« de ... à ... ») et sélection multiple avec « Indifférent » chez le/la partenaire (renseignées dans le module « + sur l'autre »).", style='List Bullet')

    # Section 3 - Demandes traitées
    add_heading_1("3. HISTORIQUE DES DEMANDES & SUIVI")

    demandes_faites = [
        "Renomme ce document en «Suivi-ARP001»",
        "Dans les types de questions remplace le type MULTI par M et corrige la signification de V A D P comme suit : (V) Le Vécu (Passé), (A) Actuel (Présent), (D) Découverte ou Poursuite (Futur), (P) Partage (Chez la personne qui partage votre quotidien ou chez les autres)",
        "La distance kilométrique est informative et ne rentre pas dans le calcul d’affinité.",
        "Les nouveaux profils ont un rôle invité et ont accès uniquement aux questions standards.",
        "Chaque profil se voit attribué un N° de profil rattaché à son pseudo et formaté comme suit : Le préfixe AFF et un N°. Il est unique. L’administrateur a un préfixe ADM. Le mien est ADM-1 pseudo Anji",
        "Dans le profil on a accès à la fiche d’identité mais aussi à des questions sur la personne. Ces questions sont définies par l’administrateur dans la banque de questions. Elles sont de niveau 0 et de classe Identité.",
        "Passe ce document en français",
        "Dans la fiche d’identité calculer l’âge en fonction de la date de naissance et vérifier la cohérence.",
        "Ajoute Pays de naissance et contrôler l’existence du pays. Remplacer Ville/région par J’habite ici (Pays, région ou département, Commune). Idem pour Je travaille là. Vérifier l’existence et la cohérence.",
        "Remplace Cible/Sexe par Sexe",
        "Remplace Style d’allure par Style",
        "À côté du bouton enregistrer, ajouter un autre bouton permettant de répondre aux questions de classe Identité avec un taux de complétude.",
        "C’est la langue de ce document word que tu dois passer en français le correcteur d’orthographe m’indique que ce n’est pas le cas.",
        "Pour Pays limite à France, UE, Hors UE. Si France limite «Région ou département» aux départements français et vérifie la cohérence avec commune renseignée.",
        "Supprime la section je travaille là",
        "Indique que la fiche est incomplète si les champs dans Pseudo, Prénom, Nom, Sexe, Date de naissance, J’habite ici et la Présentation ne sont pas renseignés. Mettre une astérisque pour indiquer que ces champs sont obligatoires.",
        "Mettre les boutons «Enregistrer» et «Répondre» en haut à droite et côte à côte. Pour enregistrer mettre uniquement l’icône correspondant à cette action. Pour Répondre mettre «+ sur vous» avec à côté un pourcentage de réponses données.",
        "Sur la fiche identité revoir le panneau du haut. On ne voit pas le nom complet, l’identifiant est sur deux lignes, réorganiser cela. Supprimer le bouton en bas «Enregistrer ma fiche».",
        "Remplacer Statut relationnel par situation de famille avec choix dans liste que tu rempliras.",
        "Ajouter «À la recherche de» avec comme choix «Échanges et Amitié», «Recherche d’un(e) partenaire», «Plus si affinité», «Je ne sais pas vraiment». Champ obligatoire.",
        "Enlève le texte sous Fiche complète «Tous les champs obligatoires» et laisse uniquement Fiche complète si 100% des réponses dans + sur vous sinon Fiche incomplète avec une bulle indiquant que pour être complète l’on doit répondre aux questions dans + sur vous.",
        "Enlever Morphologie, mensuration, allure, style & origines et en faire des questions dans + sur vous. Bien détailler les questions. Exemple une question pour le tour de poitrine, une question pour le tour de taille… Mettre les questions avec le bon classement dans la banque de questions.",
        "Remplacer NOM COMPLET par NOM et sur une ligne le nom et prénom. A la ligne l’identifiant avec le bouton copié. En dessous une petite disquette et le + sur vous avec le % (Veillez à ce que cela ne recouvre rien). A la ligne le pseudo public.",
        "Pour les questions de classe identité la réponse est une valeur (Exemple poids) ou un choix dans une liste (Exemple couleur de cheveux). Ajouter + sur l’autre et générer les mêmes questions où l’on définit les limites (Exemple en poids) ou on selectionne ce que l’on tolère (Exemple couleur de cheveux) en boite à cocher et ou on peut indiquer «Indiférent»",
        "Aligne «+ sur vous» et «+ sur l’autre»",
        "Comme «+ sur moi» «+ sur l’autre» et pris en compte pour la complétude de la fiche.",
        "Quand on clique sur «+ sur vous» et «+ sur l’autre» on accède aux questions.",
        "Les questions de classe «Identité» sont renseignées uniquement là pas dans la section «Questionnaire»",
        "Les réponses pour les questions de classe identité ne son ni de G ou M mes de type P pour des réponses précises.",
        "Dans les questions de classe «Identité» tu n’as pas créé les questions que l’on avait avant sur la fiche. Créer les questions dans la banque de question en classe identité sur la Morphologie, les mensurations, l’allure, style & origines et propose les choix de réponses correspondant aux questions. En faire des questions accessibles dans + sur vous. A partir de ces question créer les mêmes mais de type T pour tolérance accessibles dans «+ sur l’autre» . En réponse demander «de» «à» pour réponses de type valeur (Exemple le Poids, la taille) ou des coches sur les réponses de type choix dans une liste (Exemple couleur de cheveux).",
        "La classe 0 retrouve sa signification d’origine «Non définies» la classe «Identité» devient la classe 8",
        "Prendre en considération dans la banque de question le N_CIBLE (0 question quelque soit le sexe, 1 si homme, 2 si femme). Prendre en considération le sujet (N_SUJET)",
        "Retire ce que tu as créé dans les thématiques et sujets ce qui concerne les questions relatives à la classe identité",
        "Reprendre les questions dans la base MS ACCESS Affinity-Full.mdb",
        "Arrête de me créer des questions et des sujets fictifs. Fais-moi une interface pour l’administrateur créer, modifier ou supprimer des questions. Pour la classe 8 Identité une gestion de réponses précises pour le type P. Créer moi des questions dans la classe 8 de cette manière Thématique = Identité Sujet = Identité Classe = 8 Type = P Cible = 0 si valable pour un homme et une femme 1 si valable pour un homme 2 si valable pour une femme, créer les réponses possible selon la question. Créer des questions sur le physique, la Morphologie, les mensurations, l’allure, le style, les origines. Fais-en sortes que l’invité ou l’abonné puisse accéder à ces questions en cliquant dans sa fiche sur + sur moi.",
        "Dans mon Profil si je clique sur + sur moi je n'accède pas aux questions identité",
        "Ne pas mettre Questionnaire \"+ sur vous\" incomplet (0%) – Critères \"+ sur l'autre\" incomplets (0%) Mais Questionnaire \"+ sur vous\" incomplet (0%) - Questionnaire \"+ sur l'autre\" incomplet (0%)",
        "Dans la partie Gauche de la fiche d’identité ne pas faire apparaitre la classe 8",
        "Dans Créer nouvelle question pouvoir sélectionner la thématique et le sujet dans une liste. Selon le type de question adapter la gestion des réponses.",
        "Dans Niveaux & Classes de Questions griser Classe1-Standards car on ne peut la retirer et est toujours accordée. Ne pas faire apparaitre Classe 8 car accessible par + sur moi et ne peut être retirer.",
        "Dans Questionnaire avoir un filtre pour ne lister que les questions auxquelles on n’a pas répondu.",
        "Ne pas mettre d’accent sur la majuscule de « A caractère sexuel »",
        "Mettre à jour la base de données Access Affinity-Full.mdb pour prendre en compte les nouvelles questions et notamment la gestion des questions de classe 8 (Option B : tables relationnelles dédiées T_QUEST_IDENTITE et T_QUEST_IDENTITE_OPT)",
        "Propose moi 100 questions de classe 5 et 50 questions de classe 9 dans un jeu 3 et fais en sorte que je puisse les valider ou les supprimer.",
        "Dans la base Access, tu as une table T Quest dans laquelle il y a une colonne N Quest liée. Cette colonne permet de gérer des sous-questions par rapport à une question principale. Mettre ça en œuvre dans notre application (colonne n_quest_lie, filtres hiérarchiques, indicateurs visuels et arborescence)."
    ]

    for item in demandes_faites:
        p_item = doc.add_paragraph()
        r_prefix = p_item.add_run("Fait - 05/09/2026 : ")
        r_prefix.font.bold = True
        r_prefix.font.color.rgb = RGBColor(0x16, 0x65, 0x34) # Dark Green
        
        r_text = p_item.add_run(item)
        r_text.font.strike = True
        r_text.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    add_heading_1("4. NOUVELLES DEMANDES")
    p_new = doc.add_paragraph()
    r_new = p_new.add_run("*(Inscrivez ici vos prochaines demandes. Une fois prise en compte, l'assistant les passera en « Fait - [Date] » avec le texte barré).*")
    r_new.font.italic = True
    r_new.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)

    # Appliquer le code de langue français (fr-FR) sur l'ensemble du document pour le correcteur Word
    set_document_language(doc, 'fr-FR')

    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Suivi-ARP001.docx")
    try:
        doc.save(output_path)
        print(f"Document créé avec succès en langue fr-FR : {output_path}")
    except PermissionError:
        backup_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Suivi-ARP001-maj.docx")
        doc.save(backup_path)
        print(f"ATTENTION: Le fichier '{output_path}' est actuellement ouvert dans Microsoft Word.")
        print(f"Une version à jour a été enregistrée sous : '{backup_path}'.")
        print("Veuillez fermer 'Suivi-ARP001.docx' pour permettre l'écrasement direct.")

    # Supprimer l'ancien fichier Affinity.docx si présent
    old_docx = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Affinity.docx")
    if os.path.exists(old_docx):
        try:
            os.remove(old_docx)
            print("Ancien fichier Affinity.docx supprimé.")
        except Exception as e:
            print("Notice: impossible de supprimer l'ancien Affinity.docx:", e)

if __name__ == "__main__":
    create_suivi_document()

